import json

import httpx
from openpyxl import load_workbook

from invoice_extractor.cli import main
from invoice_extractor.llm import LlmConfig, LlmExtractor
from invoice_extractor.pipeline import run


def test_batch_finds_the_planted_problems(samples):
    result = run(sorted(samples.glob("*.pdf")))

    assert len(result.invoices) == 5
    assert result.failed == []
    found = {(i.file, i.check) for i in result.issues}
    assert found == {
        ("bluegate-INV-10511-wrong-total.pdf", "total"),
        ("bluegate-INV-10482-resent.pdf", "duplicate"),
    }


def test_cli_writes_excel_with_flagged_rows(samples, tmp_path):
    out = tmp_path / "invoices.xlsx"

    code = main([str(samples), "-o", str(out), "--engine", "rules"])

    assert code == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["Invoices", "Line items", "Issues"]
    statuses = {row[0]: row[1] for row in wb["Invoices"].iter_rows(min_row=2, values_only=True)}
    assert statuses["bluegate-INV-10511-wrong-total.pdf"] == "CHECK"
    assert statuses["kestrel-KL-2026-04417.pdf"] == "OK"
    assert wb["Line items"].max_row == 1 + 16


def test_broken_file_does_not_stop_the_batch(samples, tmp_path):
    bad = tmp_path / "broken.pdf"
    bad.write_bytes(b"not a pdf")

    result = run([bad, samples / "kestrel-KL-2026-04417.pdf"])

    assert len(result.invoices) == 1
    assert result.failed[0][0] == "broken.pdf"


def test_llm_path_sends_schema_and_parses_reply(samples):
    reply = {
        "vendor_name": "Kestrel Logistics GmbH",
        "vendor_tax_id": "DE312456789",
        "invoice_number": "KL-2026-04417",
        "invoice_date": "2026-09-22",
        "due_date": "2026-10-22",
        "currency": "EUR",
        "line_items": [{"description": "Sea freight", "quantity": 1, "unit_price": 2510, "amount": 2510}],
        "subtotal": 2510,
        "tax": 476.9,
        "total": 2986.9,
    }
    seen = {}

    def handler(request: httpx.Request) -> httpx.Response:
        seen["body"] = json.loads(request.content)
        seen["auth"] = request.headers["authorization"]
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(reply)}}]})

    llm = LlmExtractor(LlmConfig(api_key="test-key"), client=httpx.Client(transport=httpx.MockTransport(handler)))

    result = run([samples / "kestrel-KL-2026-04417.pdf"], llm)

    assert seen["auth"] == "Bearer test-key"
    assert seen["body"]["response_format"]["json_schema"]["strict"] is True
    assert "Kestrel Logistics" in seen["body"]["messages"][1]["content"]
    assert result.invoices[0].total == result.invoices[0].subtotal + result.invoices[0].tax
    assert result.issues == []


def test_llm_retries_on_rate_limit(monkeypatch):
    monkeypatch.setattr("invoice_extractor.llm.time.sleep", lambda _: None)
    calls = []

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(1)
        if len(calls) < 3:
            return httpx.Response(429)
        return httpx.Response(200, json={"choices": [{"message": {"content": "{}"}}]})

    llm = LlmExtractor(LlmConfig(api_key="k"), client=httpx.Client(transport=httpx.MockTransport(handler)))

    assert llm.extract("text") == {}
    assert len(calls) == 3
