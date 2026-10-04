from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .pipeline import BatchResult

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="305496")
FLAG_FILL = PatternFill("solid", fgColor="FCE4D6")
MONEY = "#,##0.000"


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], money_cols: set[int] = frozenset()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for row in rows:
        ws.append(row)
    for col in money_cols:
        for (cell,) in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell.number_format = MONEY
    for i, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def to_excel(result: BatchResult, path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    flagged = {i.file for i in result.issues if i.severity == "error"}

    ws = _sheet(
        wb,
        "Invoices",
        ["File", "Status", "Vendor", "Tax ID", "Invoice #", "Date", "Due", "Currency", "Subtotal", "Tax", "Total"],
        [
            [
                inv.source_file,
                "CHECK" if inv.source_file in flagged else "OK",
                inv.vendor_name,
                inv.vendor_tax_id,
                inv.invoice_number,
                inv.invoice_date,
                inv.due_date,
                inv.currency,
                float(inv.subtotal),
                float(inv.tax),
                float(inv.total),
            ]
            for inv in result.invoices
        ],
        money_cols={9, 10, 11},
    )
    for row in ws.iter_rows(min_row=2):
        if row[1].value == "CHECK":
            for cell in row:
                cell.fill = FLAG_FILL

    _sheet(
        wb,
        "Line items",
        ["File", "Invoice #", "Description", "Qty", "Unit price", "Amount"],
        [
            [inv.source_file, inv.invoice_number, li.description, float(li.quantity), float(li.unit_price), float(li.amount)]
            for inv in result.invoices
            for li in inv.line_items
        ],
        money_cols={5, 6},
    )

    _sheet(
        wb,
        "Issues",
        ["File", "Severity", "Check", "Message"],
        [[i.file, i.severity, i.check, i.message] for i in result.issues]
        + [[name, "error", "not_processed", reason] for name, reason in result.failed],
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def to_json(result: BatchResult, path: Path) -> Path:
    data = {
        "invoices": [inv.model_dump(mode="json") for inv in result.invoices],
        "issues": [i.__dict__ for i in result.issues],
        "failed": [{"file": f, "reason": r} for f, r in result.failed],
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2), encoding="utf-8")
    return path
